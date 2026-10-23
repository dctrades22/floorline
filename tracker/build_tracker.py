#!/usr/bin/env python3
"""Generate the Floorline tracker workbook from tracker/FORMULAS.md (spec v0.2.0) and tracker/presets_tab.csv.
No macros. Formulas are the FORMULAS.md cell formulas with PV/HIT/CENT_UP/CENT_DOWN expanded inline.

  python3 tracker/build_tracker.py            # writes tracker/Floorline-Tracker.xlsx (empty, ready to use)
  python3 tracker/build_tracker.py --demo OUT # same workbook with illustrative example rows (screenshots only)
  SOURCE_DATE_EPOCH=1792713600 python3 tracker/build_tracker.py
Bytes are reproducible. The default epoch is 2026-10-23T00:00:00Z. Set SOURCE_DATE_EPOCH
(integer UTC seconds) to use another instant. openpyxl's clock stamps are rewritten after save.
Library use: build(path, accounts=[...], logs=[...]) (used by tracker/test_tracker.py)."""
import csv, datetime, io, json, os, re, sys, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
NA, NL = 11, 1001          # last account row, last log row

# Workbook bytes must not depend on the clock. openpyxl writes created/modified and
# zip entry times from datetime.now. SOURCE_DATE_EPOCH (integer seconds, UTC) overrides
# this default. Two builds of the same tree with the same epoch are byte-identical.
BUILD_EPOCH = datetime.datetime(2026, 10, 23, 0, 0, 0)

def build_epoch():
    raw = os.environ.get("SOURCE_DATE_EPOCH")
    if not raw:
        return BUILD_EPOCH
    return datetime.datetime.fromtimestamp(int(raw), datetime.timezone.utc).replace(tzinfo=None)

def freeze_xlsx(path, epoch=None):
    """Rewrite zip entry times and core.xml created/modified to a fixed UTC instant."""
    epoch = epoch or build_epoch()
    stamp = epoch.strftime("%Y-%m-%dT%H:%M:%SZ")
    zdt = (epoch.year, epoch.month, epoch.day, epoch.hour, epoch.minute, epoch.second)
    with zipfile.ZipFile(path, "r") as zin:
        items = [(i, zin.read(i.filename)) for i in zin.infolist() if not i.is_dir()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zout:
        for info, data in sorted(items, key=lambda t: t[0].filename):
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                text = re.sub(r"(<dcterms:created[^>]*>)[^<]*", r"\g<1>" + stamp, text, count=1)
                text = re.sub(r"(<dcterms:modified[^>]*>)[^<]*", r"\g<1>" + stamp, text, count=1)
                data = text.encode("utf-8")
            ni = zipfile.ZipInfo(filename=info.filename, date_time=zdt)
            ni.compress_type = zipfile.ZIP_DEFLATED
            ni.create_system = 0
            ni.external_attr = 0
            zout.writestr(ni, data, compresslevel=6)
    with open(path, "wb") as f:
        f.write(buf.getvalue())

# ---- design tokens (mirror css/app.css) ----
BG, PANEL, INPUT, LINE, LINE2 = "0B0C0D", "121416", "1B1F23", "262A2F", "3A4048"
INK, MUT, DIM, BRAND = "ECEAE4", "8D939B", "5D636B", "FF5B1F"
GREEN, AMBER, RED, BLUE, DRED = "35D07F", "FFB224", "FF4A4A", "7CC4FF", "7A0E0E"
MONO = "Consolas"
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
F_BG, F_PANEL, F_IN = fill(BG), fill(PANEL), fill(INPUT)
thin = Side(style="thin", color=LINE); thin2 = Side(style="thin", color=LINE2)
B_CELL = Border(bottom=thin, right=thin); B_IN = Border(top=thin2, bottom=thin2, left=thin2, right=thin2)
B_HEAD = Border(bottom=Side(style="medium", color=BRAND))

# ---- notation from FORMULAS.md, expanded inline ----
PRESET_COL = {"A": "A", "B": "B", "C": "C", "D": "D", "E": "E", "F": "F", "G": "G", "H": "H", "I": "I", "J": "J",
              "K": "K", "L": "L", "M": "M", "N": "N", "O": "O", "P": "P"}
def PV(col, r): return f'IFERROR(INDEX(Presets!${col}$2:${col}$200,MATCH($B{r},Presets!$A$2:$A$200,0))&"","")'
def HIT(v, f, touch): return f'IF({touch},ROUND({v}-({f}),6)<=0,ROUND({v}-({f}),6)<0)'
def CENT_UP(x): return f'-INT(ROUND(-({x})*100,6))/100'
def CENT_DOWN(x): return f'INT(ROUND(({x})*100,6))/100'
def ISO(s):   # M-4: parse yyyy-mm-dd without locale-dependent VALUE(date text)
    return f'DATE(VALUE(LEFT({s},4)),VALUE(MID({s},6,2)),VALUE(MID({s},9,2)))'
def G(f): return f'=IFERROR({f},"INVALID")'   # final guard (FORMULAS §4)
LOGR = "'Daily Log'!$A$2:$A$1001"; LOGC = "'Daily Log'!$C$2:$C$1001"; LOGD = "'Daily Log'!$D$2:$D$1001"; LOGI = "'Daily Log'!$I$2:$I$1001"

def cap_bad(cell):
    # Blank is missing (caught earlier), not invalid. NONE is unrestricted. 0 is a real cap.
    return f'IF({cell}="",FALSE,IF({cell}="NONE",FALSE,IF(ISNUMBER({cell}),OR({cell}<0,{cell}<>INT({cell})),TRUE)))'

def account_formulas(r):
    pv = lambda c: PV(c, r)
    # AA is the param check. Q is the derived tick. Z is the active cap (minis for ES/NQ, micros for MES/MNQ).
    W = (f'=IF(AND(A{r}="",B{r}=""),"",'
         f'IF(OR(A{r}="",B{r}="",L{r}="",R{r}="",S{r}="",T{r}="",U{r}="",V{r}="",W{r}="",Z{r}="",M{r}="",AND(W{r}<>"NONE",N{r}="")),"INCOMPLETE",'
         f'IF(OR(COUNTIFS($A$2:$A${NA},A{r})>1,'
         f'AND(R{r}<>"static",R{r}<>"eod_trailing",R{r}<>"intraday_trailing"),'
         f'NOT(ISNUMBER(S{r})),S{r}<=0,ROUND(S{r}*100,6)<>INT(ROUND(S{r}*100,6)),'
         f'NOT(ISNUMBER(T{r})),T{r}<=0,ROUND(T{r}*100,6)<>INT(ROUND(T{r}*100,6)),'
         f'AND(U{r}<>"NONE",NOT(ISNUMBER(U{r}))),'
         f'AND(R{r}="static",U{r}<>"NONE"),'
         f'IF(ISNUMBER(U{r}),ROUND(S{r}+U{r}-(S{r}-T{r}),6)<=0,FALSE),'
         f'IF(ISNUMBER(U{r}),OR(U{r}<-100000000,U{r}>100000000),FALSE),'
         f'AND(V{r}<>TRUE,V{r}<>FALSE),'
         f'AND(W{r}<>"NONE",OR(NOT(ISNUMBER(W{r})),W{r}<=0)),'
         f'{cap_bad(f"X{r}")},{cap_bad(f"Y{r}")},'
         f'AND(L{r}<>"ES",L{r}<>"MES",L{r}<>"NQ",L{r}<>"MNQ"),'
         f'NOT(ISNUMBER(Q{r})),'
         f'NOT(ISNUMBER(M{r})),AND(W{r}<>"NONE",NOT(ISNUMBER(N{r}))),'
         f'S{r}>100000000,T{r}>100000000,M{r}<=0,M{r}>100000000,'
         f'IF(ISNUMBER(W{r}),W{r}>100000000,FALSE),IF(ISNUMBER(N{r}),OR(N{r}<=0,N{r}>100000000),FALSE)),"INVALID",'
         f'IF(COUNTIFS({LOGR},A{r},{LOGI},"INVALID")>0,"INVALID",'
         f'IF(OR(COUNTIFS({LOGR},A{r},{LOGI},"INCOMPLETE")>0,COUNTIFS({LOGI},"UNKNOWN_ACCOUNT")>0),"INCOMPLETE","OK")))))')
    W = "=IFERROR(" + W[1:] + ',"INVALID")'
    ok = f'$AA{r}<>"OK",$AA{r}'
    return {
        "C": f'={pv("B")}',
        "Q": G(f'IF(L{r}="","",IFERROR(INDEX(Contracts!$K$2:$K$5,MATCH(L{r},Contracts!$J$2:$J$5,0)),"INVALID"))'),
        "R": G(f'IF(D{r}<>"",D{r},{pv("C")})'),
        "S": G(f'IF(E{r}<>"",E{r},IF({pv("D")}<>"",VALUE({pv("D")}),""))'),
        "T": G(f'IF(F{r}<>"",F{r},IF({pv("E")}<>"",VALUE({pv("E")}),IF(AND({pv("F")}<>"",ISNUMBER(S{r})),VALUE({pv("F")})/100*S{r},"")))'),
        "U": G(f'IF(G{r}<>"",G{r},IF({pv("G")}="NONE","NONE",IF({pv("G")}<>"",VALUE({pv("G")}),"")))'),
        "V": G(f'IF(H{r}<>"",H{r},IF({pv("H")}="TRUE",TRUE,IF({pv("H")}="FALSE",FALSE,"")))'),
        "W": G(f'IF(I{r}<>"",I{r},IF({pv("I")}="NONE","NONE",IF({pv("I")}<>"",VALUE({pv("I")}),"")))'),
        "X": G(f'IF(J{r}<>"",J{r},IF({pv("J")}="NONE","NONE",IF({pv("J")}<>"",VALUE({pv("J")}),"")))'),
        "Y": G(f'IF(K{r}<>"",K{r},IF({pv("K")}="NONE","NONE",IF({pv("K")}<>"",VALUE({pv("K")}),"")))'),
        "Z": G(f'IF(L{r}="", "", IF(OR(L{r}="ES",L{r}="NQ"),X{r},IF(OR(L{r}="MES",L{r}="MNQ"),Y{r},"INVALID")))'),
        "AA": W,
        "AB": G(f'IF({ok},IF(R{r}="static",S{r},IF(R{r}="eod_trailing",MAX(S{r},_xlfn.MAXIFS({LOGC},{LOGR},A{r})),'
                f'MAX(S{r},M{r},_xlfn.MAXIFS({LOGC},{LOGR},A{r}),_xlfn.MAXIFS({LOGD},{LOGR},A{r})))))'),
        "AC": G(f'IF({ok},{CENT_UP(f"AN{r}")})'),
        "AD": G(f'IF({ok},AND(ISNUMBER(U{r}),R{r}<>"static",ROUND(AB{r}-T{r}-(S{r}+IF(ISNUMBER(U{r}),U{r},0)),6)>=0))'),
        "AE": G(f'IF({ok},OR(COUNTIFS({LOGR},A{r},{LOGI},"BREACHED")>0,{HIT(f"M{r}", f"AN{r}", f"V{r}")}))'),
        "AF": G(f'IF({ok},IF(AE{r},"BREACHED",{CENT_DOWN(f"MAX(0,M{r}-AN{r})")}))'),
        "AG": G(f'IF({ok},IF(AE{r},"BREACHED",IF(W{r}="NONE","NONE",{CENT_DOWN(f"MAX(0,W{r}-MAX(0,-(M{r}-N{r})))")})))'),
        "AH": G(f'IF({ok},IF(AE{r},"BREACHED",IF(AG{r}="NONE",AF{r},MIN(AF{r},AG{r}))))'),
        "AI": G(f'IF({ok},IF(AE{r},"BREACHED",IF(AND(AG{r}<>"NONE",AG{r}<AF{r}),"daily_loss","drawdown")))'),
        "AJ": G(f'IF(AO{r}<>"OK",AO{r},IF(AE{r},"BREACHED",{CENT_DOWN(f"MIN(O{r},AH{r})")}))'),
        "AK": G(f'IF(AO{r}<>"OK",AO{r},IF(AE{r},"BREACHED",IF(Z{r}="NONE",AP{r},MIN(AP{r},Z{r}))))'),
        "AL": G(f'IF(AO{r}<>"OK",AO{r},AND(Z{r}<>"NONE",AP{r}>Z{r}))'),
        "AM": G(f'IF(AO{r}<>"OK",AO{r},O{r}>AH{r})'),
        "AN": G(f'IF($AA{r}<>"OK","",IF(R{r}="static",S{r}-T{r},IF(U{r}="NONE",AB{r}-T{r},MIN(AB{r}-T{r},S{r}+U{r}))))'),
        "AO": G(f'IF({ok},IF(OR(O{r}="",P{r}="",Q{r}=""),"INCOMPLETE",IF(OR(NOT(ISNUMBER(O{r})),O{r}<0,O{r}>100000000,NOT(ISNUMBER(P{r})),P{r}<1,P{r}<>INT(P{r}),NOT(ISNUMBER(Q{r})),Q{r}<=0),"INVALID","OK")))'),
        "AP": G(f'IF(AO{r}<>"OK","",IF(AE{r},0,INT(ROUND(MIN(O{r},AH{r})/(P{r}*Q{r}),9))))'),
        "AQ": G(f'IF(AO{r}<>"OK",AO{r},IF(AE{r},"BREACHED",{CENT_DOWN(f"AK{r}*P{r}*Q{r}")}))'),
        "AR": G(f'IF(AND(A{r}="",B{r}=""),"",_xlfn.TEXTJOIN("; ",TRUE,IF(C{r}="partial","PARTIAL PRESET: verify "&{pv("L")},""),'
                f'IF(C{r}="manual","MANUAL: all params user-entered",""),'
                f'IF(OR(AND(D{r}<>"",{pv("C")}<>""),AND(E{r}<>"",{pv("D")}<>""),AND(F{r}<>"",{pv("E")}<>""),AND(G{r}<>"",{pv("G")}<>""),'
                f'AND(H{r}<>"",{pv("H")}<>""),AND(I{r}<>"",{pv("I")}<>""),AND(J{r}<>"",{pv("J")}<>""),AND(K{r}<>"",{pv("K")}<>"")),"OVERRIDES PRESET",""),'
                f'IF(AND(A{r}<>"",L{r}=""),"INSTRUMENT MISSING: choose ES, MES, NQ or MNQ",""),'
                f'IF(AND(A{r}<>"",L{r}<>"",Z{r}=""),"MAX CONTRACTS MISSING: enter the "&IF(OR(L{r}="MES",L{r}="MNQ"),"micro","mini")&" cap (number, 0, or NONE)","") ,'
                f'IF(AND({pv("N")}<>"",NOT(ISNUMBER($B$13))),"AS-OF DATE MISSING: enter today in B13 (stale check off)",""),'
                f'IF(AND({pv("N")}<>"",ISNUMBER($B$13)),IF($B$13-{ISO(pv("N"))}>30,"STALE PRESET: checked "&{pv("N")}&", over 30 days before As of",""),""),'
                f'IF({pv("M")}<>"","source: "&{pv("M")}&" checked "&{pv("N")},"")))'),
    }

def log_formulas(r):
    acc = lambda c: f'INDEX(Accounts!${c}$2:${c}${NA},$F{r})'
    t, S, DD, X, touch = acc("R"), acc("S"), acc("T"), acc("U"), acc("V")
    prior = f'_xlfn.MAXIFS($C$2:$C${NL},$A$2:$A${NL},$A{r},$B$2:$B${NL},"<"&$B{r})'
    highs = f'_xlfn.MAXIFS($D$2:$D${NL},$A$2:$A${NL},$A{r},$B$2:$B${NL},"<="&$B{r})'
    Gf = (f'=IFERROR(IF(OR($A{r}="",$F{r}=""),"",IF({t}="static",{S},IF({t}="eod_trailing",MAX({S},{prior}),'
          f'MAX({S},{prior},{highs},$C{r})))),"")')
    Hf = f'=IFERROR(IF(G{r}="","",IF({t}="static",{S}-{DD},IF({X}="NONE",G{r}-{DD},MIN(G{r}-{DD},{S}+{X})))),"")'
    dp = lambda c: f'IF({c}{r}="",FALSE,IF(ISNUMBER({c}{r}),OR({c}{r}<=0,{c}{r}>100000000,ROUND({c}{r}*100,6)<>INT(ROUND({c}{r}*100,6))),TRUE))'
    invalid = (f'OR(NOT(ISNUMBER($B{r})),NOT(ISNUMBER($C{r})),{dp("C")},{dp("D")},{dp("E")},'
               f'AND($D{r}<>"",$D{r}<MAX($C{r},$E{r})),AND($E{r}<>"",$E{r}>MIN($C{r},IF($D{r}="",$C{r},$D{r}))),'
               f'COUNTIFS($A$2:$A${NL},$A{r},$B$2:$B${NL},$B{r})>1)')
    lowv = f'IF($E{r}="",$C{r},MIN($C{r},$E{r}))'
    # spec: IF(COUNTA(A:E)=0,"",...); COUNTA is outside the §0 allowed list, so the same test is written with OR(<>"")
    If = (f'=IF(AND($A{r}="",$B{r}="",$C{r}="",$D{r}="",$E{r}=""),"",IF(OR($A{r}="",$F{r}=""),"UNKNOWN_ACCOUNT",IF(OR($B{r}="",$C{r}=""),"INCOMPLETE",IFERROR(IF({invalid},"INVALID",'
          f'IFERROR(IF({HIT(lowv, f"H{r}", touch)},"BREACHED","OK"),"INCOMPLETE")),"INVALID"))))')
    return {"F": f'=IFERROR(MATCH($A{r},Accounts!$A$2:$A${NA},0),"")', "G": Gf, "H": Hf, "I": If,
            "J": f'=IF(OR(I{r}="OK",I{r}="BREACHED"),{CENT_UP(f"H{r}")},I{r})'}

ACC_HEAD = [("A", "Account ID", 13), ("B", "Preset", 46), ("C", "Preset status", 11), ("D", "Type override", 17),
    ("E", "Start override", 12), ("F", "DD $ override", 11), ("G", "Lock offset override", 11), ("H", "Touch = breach override", 10),
    ("I", "DLL $ override", 10), ("J", "Max minis override", 10), ("K", "Max micros override", 10),
    ("L", "Instrument", 12), ("M", "Equity now (incl. open P&L)", 13), ("N", "Day reference balance", 13),
    ("O", "Risk $ / trade", 10), ("P", "Stop ticks", 7),
    ("Q", "Tick value $", 10), ("R", "Type", 18), ("S", "Start", 12), ("T", "Max DD $", 10), ("U", "Lock offset", 9),
    ("V", "Touch = breach", 8), ("W", "DLL $", 9), ("X", "Max minis", 10), ("Y", "Max micros", 10), ("Z", "Active cap", 10),
    ("AA", "Check", 12), ("AB", "Peak", 12), ("AC", "FLOOR", 14), ("AD", "Locked", 9), ("AE", "Breached", 10),
    ("AF", "DD room", 12), ("AG", "Daily room", 12), ("AH", "BINDING ROOM", 14), ("AI", "Binding", 11),
    ("AJ", "Effective risk", 11), ("AK", "CONTRACTS", 11), ("AL", "Capped", 8), ("AM", "Room-limited", 9),
    ("AN", "floor raw", 4), ("AO", "sizing check", 4), ("AP", "contracts raw", 4), ("AQ", "Actual risk", 11), ("AR", "Flags", 70)]
LOG_HEAD = [("A", "Account ID", 13), ("B", "Trading date", 13), ("C", "EOD balance", 13), ("D", "Intraday high", 13),
    ("E", "Intraday low", 13), ("F", "acct row", 4), ("G", "Peak (this row)", 13), ("H", "floor raw", 4), ("I", "Status", 12), ("J", "FLOOR", 13)]
INPUT_ACC = set("ABDEFGHIJKLMNOP"); MONEY = "#,##0.00;[Color9]-#,##0.00"

def base_sheet(ws, tab):
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

def paint(ws, rows, cols, f=F_BG):
    for r in range(1, rows + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(r, c)
            cell.fill = f; cell.font = Font(name=MONO, size=10, color=INK)

def header(ws, spec, calc_from=None):
    ws.row_dimensions[1].height = 34
    for i, (col, name, w) in enumerate(spec, 1):
        c = ws[f"{col}1"]; c.value = name
        big = name.isupper()
        c.font = Font(name=MONO, size=9, bold=True, color=BRAND if big else (INK if col in calc_from else MUT))
        c.fill = F_PANEL if col in calc_from else F_BG
        c.alignment = Alignment(wrap_text=True, vertical="bottom", horizontal="left")
        c.border = B_HEAD
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2" if ws.title == "Accounts" else "B2"

def text_sheet(wb, title, lines, tab=DIM, width=110):
    ws = wb.create_sheet(title); base_sheet(ws, tab)
    paint(ws, len(lines) + 4, 3)
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = width; ws.column_dimensions["C"].width = 3
    for i, ln in enumerate(lines, 2):
        c = ws.cell(i, 2, ln if not isinstance(ln, tuple) else ln[1])
        kind = ln[0] if isinstance(ln, tuple) else "p"
        c.font = {"h": Font(name=MONO, size=15, bold=True, color=INK), "s": Font(name=MONO, size=10, bold=True, color=BRAND),
                  "p": Font(name=MONO, size=10, color=INK), "m": Font(name=MONO, size=10, color=MUT)}[kind]
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "h": ws.row_dimensions[i].height = 28
        elif len(str(c.value)) > width - 8: ws.row_dimensions[i].height = 14 * (len(str(c.value)) // (width - 8) + 1)
    ws.protection.sheet = True
    return ws

def build(path, accounts=(), logs=(), demo_asof=None):
    wb = Workbook()
    presets = list(csv.reader(open(os.path.join(HERE, "presets_tab.csv"), newline="")))
    keys = [r[0] for r in presets[1:]]

    # ---------------- Disclaimer (first on open) ----------------
    ws = wb.active; ws.title = "Disclaimer"; wb.remove(ws)
    text_sheet(wb, "Disclaimer", [("h", "FLOORLINE  ·  Multi-Account Tracker"), ("m", "Read this first."), "",
        ("s", "DISCLAIMER"),
        "Not financial, trading, or investment advice; educational calculation tool only.",
        "Not affiliated with, endorsed by, or sponsored by any prop firm, exchange (CME Group), or platform named.",
        "Prop firm rules change without notice. Every preset shows source_url, checked_date (2026-10-14 for v0.1) and page date; verify against the firm's current rules and your dashboard before trading. Presets are not guaranteed permanent or complete; \"partial\" presets require user input.",
        "The tracker relies on user-entered balances; it cannot see intraday order of events and can differ from the firm's real-time system. The firm's own calculation is authoritative.",
        "Spreadsheet floating-point arithmetic; values are rounded to cents conservatively.", "",
        ("m", "Spec: tracker v0.2.0 / oracle v0.2.0. No macros, no scripts, no external links or data connections.")], tab=RED)

    # ---------------- How to use ----------------
    text_sheet(wb, "How to use", [("h", "How to use"),
        ("s", "REQUIREMENTS"),
        "Microsoft Excel 2019, Excel 2021 or Microsoft 365 (uses MAXIFS and TEXTJOIN; Excel 2016 and earlier show #NAME? and are NOT supported).",
        "Tested: formulas were recalculated and checked in LibreOffice only. Excel and Google Sheets have NOT been tested; treat both as unverified until you confirm the results yourself.",
        ("s", "GOOGLE SHEETS IMPORT (UNTESTED)"),
        "1. Open drive.google.com, click New > File upload and choose Floorline-Tracker.xlsx.",
        "2. Right-click the uploaded file > Open with > Google Sheets (or set Drive Settings > 'Convert uploads' before uploading).",
        "3. In Sheets: File > Save as Google Sheets if it opened in Office-compatibility mode, then work in the converted copy.",
        "4. Check the Accounts tab: Check column shows OK for filled rows and no cell shows #NAME? or #ERROR!. Compare a few results with your firm dashboard before relying on it.",
        "5. Not verified in Sheets: whether sheet protection, dropdowns, data validation and colour rules survive the import. Pasted values can bypass dropdowns and validation; the formulas check many inputs, but not every possible mistake.",
        ("s", "DAILY ROUTINE"),
        "1. Accounts: one row per account (max 10). Enter Account ID and pick a Preset (or Manual). Read the Flags column.",
        "2. Fill any field named in 'PARTIAL PRESET: verify ...' using the override columns D-K (amber flag until resolved).",
        "3. Choose Instrument (ES, MES, NQ or MNQ). Tick value fills in from that instrument (ES 12.50, MES 1.25, NQ 5.00, MNQ 0.50) and cannot be typed. Enter Equity now, Day reference balance, Risk $ and Stop ticks. ES/NQ use the mini cap; MES/MNQ use the micro cap. Blank required cap = INCOMPLETE. 0 stays 0. NONE removes only the contract cap. Risk, trailing drawdown and the daily loss limit stay active.",
        "4. Daily Log: after each session add one row per account: Trading date (firm session date), EOD balance, intraday high and intraday low.",
        "5. Read FLOOR, BINDING ROOM and CONTRACTS on Accounts. Then tick the Checklist. If your firm dashboard disagrees, the firm's number wins.",
        "6. Enter today's date in Accounts!B13 (As of). Flags shows AS-OF DATE MISSING while it is blank and STALE PRESET when a preset was checked more than 30 days before it. The workbook never uses TODAY().",
        "7. If Accounts!B14 shows a red Log check message, a Daily Log row has an unknown or blank account ID; every account shows INCOMPLETE until it is fixed.",
        ("s", "STATUS WORDS"),
        "OK / a number = calculated.  INVALID = a value or rule is wrong (money must be 0 < x <= 100,000,000 with at most 2 decimals; instrument must be ES, MES, NQ or MNQ).  INCOMPLETE = a required input is missing, incl. a blank instrument or a blank cap for that instrument (type NONE for no contract cap), an incomplete log row, or any unknown-account log row.  UNKNOWN_ACCOUNT = log row whose account ID matches no account.  BREACHED = the floor was hit; rooms and size are blocked.  NONE = no daily loss limit, or no contract cap. Risk, trailing drawdown and daily loss stay active either way.",
        ("s", "COLOURS"),
        "Green = buffer above 50% of max drawdown or lock reached.  Amber = 25-50% buffer, partial preset, overrides, stale preset.  Red = under 25%, INVALID or BREACHED.  Blue = sizing was capped or clamped to the room.",
        ("s", "LIMITS"),
        "Intraday trailing: a daily row cannot encode whether the low came before the high; the tracker assumes the high came first (higher, more conservative floor). Percent daily limits and payout/withdrawal floor locks must be entered manually. Futures only. 10 accounts, 1,000 log rows.",
        "Matching ignores letter case: account IDs must be unique ignoring case (p16 and P16 are duplicates), log row p16 belongs to account P16, and none / None count as NONE. Spaces are NOT ignored: 'Static ' with a trailing space is INVALID."], tab=BLUE)

    # ---------------- Accounts ----------------
    ws = wb.create_sheet("Accounts"); base_sheet(ws, BRAND)
    paint(ws, 16, len(ACC_HEAD))
    calc = {c for c, _, _ in ACC_HEAD if c not in INPUT_ACC}
    header(ws, ACC_HEAD, calc)
    for r in range(2, NA + 1):
        ws.row_dimensions[r].height = 20
        for col, f in account_formulas(r).items(): ws[f"{col}{r}"] = f
        for col, _, _ in ACC_HEAD:
            c = ws[f"{col}{r}"]
            inp = col in INPUT_ACC
            c.fill = F_IN if inp else F_BG
            c.border = B_IN if inp else B_CELL
            c.protection = Protection(locked=not inp)
            c.alignment = Alignment(vertical="center", horizontal="left" if col in ("A", "B", "L", "AA", "R", "AI", "AR") else "right")
            bold = col in ("AC", "AH", "AK")
            muted = col in ("C", "R", "S", "T", "U", "V", "W", "X", "Y", "Z")
            c.font = Font(name=MONO, size=12 if bold else 10, bold=bold, color=INK if (inp or bold) else MUT if muted else INK)
            if col in ("E", "F", "I", "M", "N", "O", "S", "T", "W", "AB", "AC", "AF", "AG", "AH", "AJ", "AQ"): c.number_format = MONEY
            if col == "Q": c.number_format = "0.00"
    ws["A13"] = "As of"; ws["A13"].font = Font(name=MONO, size=9, bold=True, color=MUT)
    ws["B13"].fill = F_IN; ws["B13"].border = B_IN; ws["B13"].number_format = "yyyy-mm-dd"; ws["B13"].protection = Protection(locked=False)
    ws["B13"].font = Font(name=MONO, size=10, color=INK); ws["B13"].alignment = Alignment(horizontal="left")
    ws["C13"] = "← today's date (enables the 30-day stale-preset check; no TODAY() used)"; ws["C13"].font = Font(name=MONO, size=9, color=DIM)
    ws["A14"] = "Log check"; ws["A14"].font = Font(name=MONO, size=9, bold=True, color=MUT)
    ws["B14"] = (f'=IF(COUNTIFS({LOGI},"UNKNOWN_ACCOUNT")>0,COUNTIFS({LOGI},"UNKNOWN_ACCOUNT")&'
                 '" Daily Log row(s) with unknown account ID - fix before trading","")')
    ws["B14"].font = Font(name=MONO, size=10, bold=True, color=RED)
    ws.conditional_formatting.add("B14:C14", FormulaRule(formula=['$B$14<>""'], fill=fill(RED), font=Font(color="000000", bold=True)))
    ws["A15"] = "Inputs: shaded cells except C and Q.  Instrument L is ES, MES, NQ or MNQ.  Tick value Q is calculated from the instrument and is not typed.  ES/NQ use the mini cap; MES/MNQ use the micro cap.  FLOOR / BINDING ROOM / CONTRACTS are the numbers to trade by."
    ws["A15"].font = Font(name=MONO, size=9, color=DIM)
    for col in ("AN", "AO", "AP"): ws.column_dimensions[col].hidden = True

    rng = f"$A$2:$AR${NA}"
    def cf(ref, formula, **kw): ws.conditional_formatting.add(ref, FormulaRule(formula=[formula], stopIfTrue=False, **kw))
    font = lambda c, b=False: Font(color=c, bold=b)
    # whole-row BREACHED, then cell-level states
    cf(rng, f'$AE2=TRUE', fill=fill(DRED), font=font("FFFFFF", True))
    cf(f"AA2:AA{NA}", 'AA2="INVALID"', fill=fill(RED), font=font("000000", True))
    cf(f"AA2:AA{NA}", 'AA2="INCOMPLETE"', fill=fill(AMBER), font=font("000000", True))
    cf(f"AA2:AA{NA}", 'AA2="OK"', font=font(GREEN, True))
    cf(f"AD2:AD{NA}", 'AD2=TRUE', font=font(GREEN, True))
    cf(f"AF2:AF{NA}", 'AND(ISNUMBER(AF2),ISNUMBER(T2),AF2<0.25*T2)', font=font("FF8A3D", True))
    lvl = f'AND(ISNUMBER($AH2),ISNUMBER($T2))'
    cf(f"AH2:AH{NA}", f'AND({lvl},$AH2<=0.25*$T2)', font=font(RED, True))
    cf(f"AH2:AH{NA}", f'AND({lvl},$AH2<=0.5*$T2)', font=font(AMBER, True))
    cf(f"AH2:AH{NA}", f'{lvl}', font=font(GREEN, True))
    cf(f"AC2:AC{NA}", f'{lvl}', font=Font(color=INK, bold=True))
    cf(f"AL2:AM{NA}", 'AL2=TRUE', fill=fill("0E2A40"), font=font(BLUE, True))
    cf(f"AK2:AK{NA}", 'OR($AL2=TRUE,$AM2=TRUE)', font=font(BLUE, True))
    cf(f"AR2:AR{NA}", 'OR(ISNUMBER(SEARCH("PARTIAL",AR2)),ISNUMBER(SEARCH("OVERRIDES",AR2)),ISNUMBER(SEARCH("MISSING",AR2)),ISNUMBER(SEARCH("STALE",AR2)))', font=font(AMBER))
    cf(f"C2:C{NA}", 'C2="verified"', font=font(GREEN, True))
    cf(f"C2:C{NA}", 'C2="partial"', font=font(AMBER, True))
    cf(f"B2:C{NA}", f'ISNUMBER(SEARCH("STALE PRESET",$AR2))', fill=fill("3A2A08"), font=font(AMBER, True))

    dv = lambda **k: DataValidation(showErrorMessage=True, errorStyle="stop", **k)
    v = dv(type="list", formula1="Presets!$A$2:$A$200", allow_blank=True); v.add(f"B2:B{NA}"); ws.add_data_validation(v)
    v = dv(type="list", formula1='"static,eod_trailing,intraday_trailing"', allow_blank=True); v.add(f"D2:D{NA}"); ws.add_data_validation(v)
    v = dv(type="list", formula1='"TRUE,FALSE"', allow_blank=True); v.add(f"H2:H{NA}"); ws.add_data_validation(v)
    money = lambda ref, extra="TRUE": f'AND(ISNUMBER({ref}),ROUND({ref}*100,6)=INT(ROUND({ref}*100,6)),{extra})'
    for col, extra in (("E", "E2>0,E2<=100000000"), ("F", "F2>0,F2<=100000000"), ("M", "M2>0,M2<=100000000"), ("N", "N2>0,N2<=100000000"), ("O", "O2>=0,O2<=100000000")):
        v = dv(type="custom", formula1=money(f"{col}2", extra), allow_blank=True, error="Number with at most 2 decimals, up to 100,000,000"); v.add(f"{col}2:{col}{NA}"); ws.add_data_validation(v)
    v = dv(type="custom", formula1=f'OR(G2="NONE",{money("G2", "G2>=-100000000,G2<=100000000")})', allow_blank=True, error="Number (2 dp, may be <=0) or NONE"); v.add(f"G2:G{NA}"); ws.add_data_validation(v)
    v = dv(type="custom", formula1=f'OR(I2="NONE",{money("I2", "I2>0,I2<=100000000")})', allow_blank=True, error="Dollar amount > 0 or NONE"); v.add(f"I2:I{NA}"); ws.add_data_validation(v)
    for col in ("J", "K"):
        v = dv(type="custom", formula1=f'OR({col}2="NONE",AND(ISNUMBER({col}2),{col}2>=0,{col}2=INT({col}2)))', allow_blank=True, error="Whole number >= 0 (0 is allowed) or NONE"); v.add(f"{col}2:{col}{NA}"); ws.add_data_validation(v)
    v = dv(type="list", formula1="Contracts!$J$2:$J$5", allow_blank=True, error="Choose ES, MES, NQ or MNQ"); v.add(f"L2:L{NA}"); ws.add_data_validation(v)
    v = dv(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True, error="Whole number >= 1"); v.add(f"P2:P{NA}"); ws.add_data_validation(v)
    v = dv(type="date", operator="greaterThan", formula1="1", allow_blank=True); v.add("B13"); ws.add_data_validation(v)
    ws.protection.sheet = True; ws.protection.formatColumns = False; ws.protection.formatRows = False
    for i, a in enumerate(accounts, 2):
        for col, val in a.items(): ws[f"{col}{i}"] = val
    if demo_asof: ws["B13"] = demo_asof

    # ---------------- Daily Log ----------------
    ws = wb.create_sheet("Daily Log"); base_sheet(ws, INK)
    lcalc = {"F", "G", "H", "I", "J"}
    header(ws, LOG_HEAD, lcalc)
    for r in range(2, NL + 1):
        ws.row_dimensions[r].height = 18
        for col, f in log_formulas(r).items(): ws[f"{col}{r}"] = f
        for col, _, _ in LOG_HEAD:
            c = ws[f"{col}{r}"]; inp = col not in lcalc
            c.fill = F_IN if inp else F_BG; c.border = B_IN if inp else B_CELL
            c.protection = Protection(locked=not inp)
            c.font = Font(name=MONO, size=10, bold=col == "J", color=INK if inp or col in "IJ" else MUT)
            c.alignment = Alignment(horizontal="left" if col in "AI" else "right", vertical="center")
            if col in "CDEGJ": c.number_format = MONEY
            if col == "B": c.number_format = "yyyy-mm-dd"
    for col in ("F", "H"): ws.column_dimensions[col].hidden = True
    lr = f"$A$2:$J${NL}"
    ws.conditional_formatting.add(lr, FormulaRule(formula=['$I2="BREACHED"'], fill=fill(DRED), font=Font(color="FFFFFF", bold=True)))
    ws.conditional_formatting.add(f"A2:J{NL}", FormulaRule(formula=['$I2="UNKNOWN_ACCOUNT"'], fill=fill(RED), font=Font(color="000000", bold=True)))
    ws.conditional_formatting.add(f"I2:J{NL}", FormulaRule(formula=['$I2="INVALID"'], fill=fill(RED), font=Font(color="000000", bold=True)))
    ws.conditional_formatting.add(f"I2:J{NL}", FormulaRule(formula=['$I2="INCOMPLETE"'], fill=fill(AMBER), font=Font(color="000000", bold=True)))
    ws.conditional_formatting.add(f"I2:I{NL}", FormulaRule(formula=['$I2="OK"'], font=Font(color=GREEN, bold=True)))
    v = dv(type="list", formula1=f"Accounts!$A$2:$A${NA}", allow_blank=True); v.add(f"A2:A{NL}"); ws.add_data_validation(v)
    v = dv(type="date", operator="greaterThan", formula1="1", allow_blank=True, error="Enter a valid date"); v.add(f"B2:B{NL}"); ws.add_data_validation(v)
    for col in "CDE":
        v = dv(type="custom", formula1=money(f"{col}2", f"{col}2>0,{col}2<=100000000"), allow_blank=True, error="Number with at most 2 decimals, up to 100,000,000"); v.add(f"{col}2:{col}{NL}"); ws.add_data_validation(v)
    ws.protection.sheet = True
    for i, row in enumerate(logs, 2):
        for col, val in row.items(): ws[f"{col}{i}"] = val

    # ---------------- Presets ----------------
    ws = wb.create_sheet("Presets"); base_sheet(ws, GREEN)
    nr, nc = len(presets), len(presets[0])
    paint(ws, nr, nc)
    numeric = {"start_balance", "dd_amount", "dd_pct", "max_minis", "max_micros"}
    hdr = presets[0]
    for r, row in enumerate(presets, 1):
        for c, val in enumerate(row, 1):
            cell = ws.cell(r, c)
            # keep tokens as text so PV()'s &"" / VALUE() path works identically in Excel and Sheets
            cell.value = val if val != "" else None
            if r > 1: cell.number_format = "@"
            if r == 1: cell.font = Font(name=MONO, size=9, bold=True, color=MUT); cell.border = B_HEAD
            elif c == 2: cell.font = Font(name=MONO, size=10, bold=True, color=GREEN if val == "verified" else AMBER if val == "partial" else MUT)
    for c, w in zip("ABCDEFGHIJKLMNOP", (52, 9, 16, 11, 9, 7, 8, 7, 8, 7, 7, 40, 60, 11, 11, 7)): ws.column_dimensions[c].width = w
    ws.freeze_panes = "B2"; ws.protection.sheet = True

    # ---------------- Contracts ----------------
    cj = json.load(open(os.path.join(ROOT, "data", "contracts.json")))
    ws = wb.create_sheet("Contracts"); base_sheet(ws, MUT)
    head = ["symbol", "name", "tick_size", "tick_value_usd", "status", "checked_date", "source_url", "note"]
    paint(ws, max(len(cj["contracts"]) + 3, 6), 11)
    for c, h in enumerate(head, 1):
        x = ws.cell(1, c, h); x.font = Font(name=MONO, size=9, bold=True, color=MUT); x.border = B_HEAD
    for r, k in enumerate(cj["contracts"], 2):
        for c, h in enumerate(head, 1):
            x = ws.cell(r, c, k[h]); x.number_format = "@"
            if h == "status": x.font = Font(name=MONO, size=10, bold=True, color=GREEN if k[h] == "verified" else AMBER)
            if h == "symbol": x.font = Font(name=MONO, size=10, bold=True, color=INK)
    # Locked lookup for Accounts tick formulas. Dropdown and tick share this list, so a mismatch cannot be typed.
    expect = {"ES": 12.5, "MES": 1.25, "NQ": 5.0, "MNQ": 0.5}
    by_sym = {c["symbol"]: float(c["tick_value_usd"]) for c in cj["contracts"]}
    if any(by_sym.get(s) != v for s, v in expect.items()):
        raise SystemExit(f"contracts.json tick values disagree with the instrument table: {by_sym}")
    ws["J1"] = "instrument"; ws["K1"] = "tick_value"
    ws["J1"].font = Font(name=MONO, size=9, bold=True, color=BRAND); ws["K1"].font = Font(name=MONO, size=9, bold=True, color=BRAND)
    for i, sym in enumerate(("ES", "MES", "NQ", "MNQ"), 2):
        ws.cell(i, 10, sym).font = Font(name=MONO, size=10, bold=True, color=INK)
        cell = ws.cell(i, 11, expect[sym]); cell.number_format = "0.00"; cell.font = Font(name=MONO, size=10, color=INK)
    ws.cell(len(cj["contracts"]) + 3, 1, "Accounts tick value looks up J2:K5 only (ES 12.50, MES 1.25, NQ 5.00, MNQ 0.50). Those four match CME Rulebook chapters 358, 353, 359 and 361, checked 2026-10-23. Other rows are reference and may still be verified indirectly. Confirm on cmegroup.com.").font = Font(name=MONO, size=9, color=AMBER)
    for c, w in zip("ABCDEFGHJK", (8, 28, 10, 13, 17, 12, 50, 80, 12, 12)): ws.column_dimensions[c].width = w
    ws.freeze_panes = "B2"; ws.protection.sheet = True

    # ---------------- Checklist ----------------
    ws = wb.create_sheet("Checklist", index=3); base_sheet(ws, AMBER)
    items = ["Preset matches the plan purchased (firm, plan, size, purchase date — older purchases may have legacy rules).",
        "Opened source_url; rule text still matches; checked_date within 30 days.",
        "All needs_user / PARTIAL flags resolved.",
        "Lock offset confirmed (0 = start, X = start+X, NONE).",
        "Touch rule confirmed (\"hit/touch\" = TRUE, \"below\" = FALSE).",
        "DLL entered in $ for today's reset (% plans: firm dashboard value).",
        "Day reference balance = firm's reset reference (incl./excl. unrealized per firm).",
        "Mini cap and micro cap confirmed. ES/NQ use max minis; MES/MNQ use max micros. Blank required cap is incomplete. NONE removes only the contract cap. 0 means zero contracts.",
        "Today's Daily Log row entered incl. intraday low.",
        "Instrument is ES, MES, NQ or MNQ. Tick value is calculated (ES 12.50, MES 1.25, NQ 5.00, MNQ 0.50) and is not typed.",
        "Firm dashboard floor equals tracker floor; if not, the firm's number wins."]
    paint(ws, len(items) + 8, 5)
    ws["A1"] = "PRE-SESSION CHECKLIST"; ws["A1"].font = Font(name=MONO, size=14, bold=True, color=INK)
    ws["A2"] = "Pick the account, tick each line (type x or ✓). Done = 11/11."; ws["A2"].font = Font(name=MONO, size=9, color=MUT)
    for c, h in zip("ABCD", ("#", "Done", "Account", "Check")):
        x = ws[f"{c}4"]; x.value = h; x.font = Font(name=MONO, size=9, bold=True, color=MUT); x.border = B_HEAD
    for i, t in enumerate(items, 5):
        ws[f"A{i}"] = i - 4; ws[f"A{i}"].font = Font(name=MONO, size=10, bold=True, color=BRAND); ws[f"A{i}"].alignment = Alignment(horizontal="center", vertical="center")
        for col in "BC":
            x = ws[f"{col}{i}"]; x.fill = F_IN; x.border = B_IN; x.protection = Protection(locked=False)
            x.alignment = Alignment(horizontal="center", vertical="center"); x.font = Font(name=MONO, size=12, bold=True, color=GREEN)
        ws[f"C{i}"].font = Font(name=MONO, size=10, color=INK)
        ws[f"D{i}"] = t; ws[f"D{i}"].font = Font(name=MONO, size=10, color=INK)
        ws[f"D{i}"].alignment = Alignment(wrap_text=True, vertical="center"); ws.row_dimensions[i].height = 30
    n = 5 + len(items) + 1
    ws[f"C{n}"] = "Done"; ws[f"C{n}"].font = Font(name=MONO, size=9, bold=True, color=MUT); ws[f"C{n}"].alignment = Alignment(horizontal="right")
    ws[f"D{n}"] = f'=COUNTIFS(B5:B{4+len(items)},"<>")&" / {len(items)}"'
    ws[f"D{n}"].font = Font(name=MONO, size=14, bold=True, color=INK)
    ws.conditional_formatting.add(f"D{n}", FormulaRule(formula=[f'COUNTIFS(B5:B{4+len(items)},"<>")={len(items)}'], font=Font(color=GREEN, bold=True)))
    ws.conditional_formatting.add(f"A5:D{4+len(items)}", FormulaRule(formula=['$B5<>""'], font=Font(color=DIM)))
    v = dv(type="list", formula1='"✓,x"', allow_blank=True); v.add(f"B5:B{4+len(items)}"); ws.add_data_validation(v)
    v = dv(type="list", formula1=f"Accounts!$A$2:$A${NA}", allow_blank=True); v.add(f"C5:C{4+len(items)}"); ws.add_data_validation(v)
    for c, w in zip("ABCD", (5, 7, 14, 100)): ws.column_dimensions[c].width = w
    ws.protection.sheet = True

    wb.active = 0
    wb.calculation.fullCalcOnLoad = True
    epoch = build_epoch()
    wb.properties.creator = "Floorline"
    wb.properties.lastModifiedBy = "Floorline"
    wb.properties.title = "Floorline Multi-Account Tracker"
    wb.properties.created = epoch
    wb.properties.modified = epoch
    wb.properties.revision = "1"
    wb.save(path)
    freeze_xlsx(path, epoch)
    return path

DEMO_ACC = [
    {"A": "TS-50K-01", "B": "Topstep | Trading Combine | 50000 [P001]", "I": 1000, "L": "MES", "M": 51240.5, "N": 51550.5, "O": 300, "P": 12},
    {"A": "APEX-50", "B": "Apex Trader Funding | Intraday Trailing Evaluation | 50000 [P006]", "G": "NONE", "J": "NONE", "K": 40, "L": "MNQ", "M": 50410, "O": 250, "P": 16},
    {"A": "LUCID-25", "B": "Lucid Trading | LucidFlex (eval + funded) | 25000 [P018]", "I": "NONE", "J": "NONE", "K": "NONE", "L": "MES", "M": 25120, "O": 100, "P": 8},
    {"A": "TS-100K", "B": "Topstep | Trading Combine | 100000 [P002]", "I": "NONE", "L": "NQ", "M": 97050, "O": 400, "P": 10},
    {"A": "MANUAL-1", "B": "Manual", "D": "static", "E": 25000, "F": 1000, "G": "NONE", "H": True, "I": "NONE", "J": 5, "K": 50, "L": "ES", "M": 24300, "O": 200, "P": 12},
]
d = datetime.date
DEMO_LOG = [
    {"A": "TS-50K-01", "B": d(2026, 10, 12), "C": 51100, "D": 51320, "E": 50880},
    {"A": "TS-50K-01", "B": d(2026, 10, 13), "C": 51550.5, "D": 51700, "E": 51010},
    {"A": "APEX-50", "B": d(2026, 10, 12), "C": 50620, "D": 50980, "E": 50300},
    {"A": "APEX-50", "B": d(2026, 10, 13), "C": 50410, "D": 50700, "E": 49990},
    {"A": "LUCID-25", "B": d(2026, 10, 13), "C": 25120, "D": 25400, "E": 24980},
    {"A": "TS-100K", "B": d(2026, 10, 12), "C": 99800, "D": 100250, "E": 98900},
    {"A": "TS-100K", "B": d(2026, 10, 13), "C": 97050, "D": 99900, "E": 96950},
    {"A": "MANUAL-1", "B": d(2026, 10, 13), "C": 24300, "D": 24600, "E": 24150},
]

if __name__ == "__main__":
    if "--demo" in sys.argv:
        out = sys.argv[sys.argv.index("--demo") + 1]
        print("wrote", build(out, DEMO_ACC, DEMO_LOG, demo_asof=d(2026, 10, 14)))
    else:
        print("wrote", build(os.path.join(HERE, "Floorline-Tracker.xlsx")))
