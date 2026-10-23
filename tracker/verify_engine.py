#!/usr/bin/env python3
"""Inspect all 10 account rows, recalculate representative cases in LibreOffice, compare to JS and Python.

  python3 tracker/verify_engine.py REPORT.json WORKBOOK.xlsx
The workbook is a private artifact. Do not commit it.
"""
import json, os, subprocess, sys, tempfile
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(ROOT, "golden"))
import build_tracker as bt
import oracle

# One case per account row. Rows 2-5 are the required instrument results.
ROWS = [
    ("MES", 200, 12, 1, 50, {"contracts": 13, "tick": 1.25, "status": "OK"}),
    ("ES", 200, 12, 5, 1, {"contracts": 1, "tick": 12.5, "status": "OK"}),
    ("NQ", 200, 8, 10, "NONE", {"contracts": 5, "tick": 5, "status": "OK"}),
    ("MNQ", 200, 8, "NONE", 100, {"contracts": 50, "tick": 0.5, "status": "OK"}),
    ("MES", 200, 12, 5, None, {"contracts": "INCOMPLETE", "tick": 1.25, "status": "INCOMPLETE"}),
    ("MES", 200, 12, "NONE", 0, {"contracts": 0, "tick": 1.25, "status": "OK"}),
    ("MES", 10000, 12, "NONE", "NONE", {"contracts": 13, "tick": 1.25, "status": "OK", "binding": "daily_loss"}),
    ("ES", 200, 12, 5, "NONE", {"contracts": "BREACHED", "tick": 12.5, "status": "OK", "equity": 48000}),
    ("ZZ", 200, 12, 5, 50, {"contracts": "INVALID", "tick": "INVALID", "status": "INVALID"}),
    ("", 200, 12, 5, 50, {"contracts": "INCOMPLETE", "tick": None, "status": "INCOMPLETE"}),
]

def account(i, row):
    inst, risk, ticks, minis, micros, exp = row
    equity = exp.get("equity", 49000 if exp.get("binding") != "daily_loss" else 49000)
    if exp.get("binding") == "daily_loss":
        equity = 49000
    a = {"A": f"R{i}", "B": "Manual", "D": "static", "E": 50000, "F": 2000, "G": "NONE", "H": True,
         "I": 1000 if exp.get("binding") == "daily_loss" else "NONE", "M": equity, "O": risk, "P": ticks}
    if exp.get("binding") == "daily_loss":
        a["N"] = 49800
    if inst != "":
        a["L"] = inst
    if minis is not None:
        a["J"] = minis
    if micros is not None:
        a["K"] = micros
    return a

def main(report_path, book_path):
    empty = os.path.join(tempfile.mkdtemp(), "empty.xlsx")
    bt.build(empty)
    wb = load_workbook(empty)
    ws = wb["Accounts"]
    formulas = []
    for r in range(2, 12):
        formulas.append({
            "row": r,
            "tick_Q": ws[f"Q{r}"].value,
            "active_cap_Z": ws[f"Z{r}"].value,
            "contracts_AK": ws[f"AK{r}"].value,
            "check_AA": ws[f"AA{r}"].value,
            "tick_locked": bool(ws[f"Q{r}"].protection.locked),
            "instrument_unlocked": not bool(ws[f"L{r}"].protection.locked),
        })
        for col in ("Q", "Z", "AK", "AA"):
            text = ws[f"{col}{r}"].value or ""
            if f"{col[0]}{r}" not in text and str(r) not in text:
                raise SystemExit(f"row {r} {col} formula does not reference its own row")
    vals = []
    for dv in ws.data_validations.dataValidation:
        vals.append({"type": dv.type, "sqref": str(dv.sqref), "formula1": dv.formula1})
    covered = " ".join(v["sqref"] for v in vals)
    for need in ("L2:L11", "J2:J11", "K2:K11"):
        if need not in covered:
            raise SystemExit("missing validation " + need + " in " + covered)
    if any("Q2:Q11" in v["sqref"] for v in vals):
        raise SystemExit("tick value column is an input; it must not have data validation")
    accs = [account(i, row) for i, row in enumerate(ROWS, 1)]
    bt.build(book_path, accs)
    out = tempfile.mkdtemp()
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", out, book_path],
                   check=True, capture_output=True, timeout=180)
    calc = load_workbook(os.path.join(out, os.path.basename(book_path)), data_only=True)["Accounts"]
    report_rows = []
    mismatches = []
    for i, row in enumerate(ROWS):
        inst, risk, ticks, minis, micros, exp = row
        r = i + 2
        got = {
            "status": calc[f"AA{r}"].value,
            "tick": calc[f"Q{r}"].value,
            "active_cap": calc[f"Z{r}"].value,
            "contracts": calc[f"AK{r}"].value,
            "binding": calc[f"AI{r}"].value,
            "breached": calc[f"AE{r}"].value,
        }
        py = None
        if inst in ("ES", "MES", "NQ", "MNQ"):
            klass = "max_minis" if inst in ("ES", "NQ") else "max_micros"
            cap = minis if klass == "max_minis" else micros
            inp = {"instrument": inst, "risk": str(risk), "binding_room": "200" if exp.get("binding") == "daily_loss" else ("1000" if exp.get("equity") == 48000 else "10000"),
                   "stop_ticks": str(ticks)}
            if cap is not None:
                inp[klass] = str(cap)
            other = "max_micros" if klass == "max_minis" else "max_minis"
            other_v = micros if other == "max_micros" else minis
            if other_v is not None:
                inp[other] = str(other_v)
            py = oracle.evaluate("futures", inp)
        ok_status = got["status"] == exp["status"]
        ok_tick = (got["tick"] == exp["tick"]) or (exp["tick"] is None and (got["tick"] in ("", None)))
        ok_contracts = got["contracts"] == exp["contracts"]
        if exp.get("binding"):
            ok_contracts = ok_contracts and got["binding"] == exp["binding"]
        if py and py.get("ok") and exp["status"] == "OK" and exp["contracts"] != "BREACHED":
            if py["contracts"] != got["contracts"]:
                mismatches.append({"row": r, "why": "python != libreoffice", "python": py["contracts"], "sheet": got["contracts"]})
        if not (ok_status and ok_tick and ok_contracts):
            mismatches.append({"row": r, "expected": exp, "got": got, "python": py})
        report_rows.append({"row": r, "input": {"instrument": inst, "risk": risk, "ticks": ticks, "max_minis": minis, "max_micros": micros},
                            "libreoffice": got, "python": py})
    # JS for the four required rows, with the same inputs the UI would send.
    js_code = """
const Oracle = require("./js/oracle.js");
const cases = [
  {instrument:"MES", risk:"200", binding_room:"10000", stop_ticks:"12", max_minis:"1", max_micros:"50"},
  {instrument:"ES", risk:"200", binding_room:"10000", stop_ticks:"12", max_minis:"5", max_micros:"1"},
  {instrument:"NQ", risk:"200", binding_room:"10000", stop_ticks:"8", max_minis:"10", max_micros:"NONE"},
  {instrument:"MNQ", risk:"200", binding_room:"10000", stop_ticks:"8", max_minis:"NONE", max_micros:"100"}
];
process.stdout.write(JSON.stringify(cases.map(c => Oracle.evaluate("futures", c))));
"""
    js_run = subprocess.run(["node", "-e", js_code], cwd=ROOT, check=True, capture_output=True, text=True)
    js = json.loads(js_run.stdout)
    js_ok = [js[0]["contracts"], js[1]["contracts"], js[2]["contracts"], js[3]["contracts"]] == [13, 1, 5, 50]
    doc = {"engine": "LibreOffice via soffice --headless --convert-to xlsx",
           "formulas_all_10_rows": formulas, "validations": vals,
           "rows": report_rows, "js_required_contracts": js_ok, "js": js,
           "mismatches": mismatches}
    os.makedirs(os.path.dirname(os.path.abspath(report_path)), exist_ok=True)
    with open(report_path, "w") as f:
        json.dump(doc, f, indent=2, default=str)
        f.write("\n")
    print("rows", len(formulas), "validations", len(vals), "js_required", js_ok, "mismatches", len(mismatches))
    print("wrote", report_path)
    print("wrote", book_path)
    return 1 if mismatches or not js_ok else 0

if __name__ == "__main__":
    rep, book = sys.argv[1], sys.argv[2]
    sys.exit(main(rep, book))
